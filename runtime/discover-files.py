"""Bounded source discovery. Reads explicit files locally or with a private Blob SAS handout."""
import argparse, csv, datetime, decimal, hashlib, io, json, pathlib, re, sys, urllib.request, urllib.parse, zipfile
MAX_BYTES=20*1024*1024
MAX_TOTAL=100*1024*1024
class DiscoveryError(ValueError): pass
def fail(message): raise DiscoveryError(message)
def unique(pairs):
    out={}
    for k,v in pairs:
        if k in out: fail('Duplicate JSON property')
        out[k]=v
    return out
def parse_json(raw):
    return json.loads(raw,parse_float=decimal.Decimal,object_pairs_hook=unique,parse_constant=lambda _:fail('Non-finite JSON number'))
def name(value):
    if not isinstance(value,str) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,62}',value): fail('Column names require reviewed simple identifiers')
    return value

def infer(values, field, missing=0, depth=0):
    if depth>8: fail('Nested schema depth exceeds 8')
    present=[v for v in values if v is not None]
    kinds=set('boolean' if isinstance(v,bool) else 'integer' if isinstance(v,int) else 'number' if isinstance(v,(float,decimal.Decimal)) else 'timestamp' if isinstance(v,datetime.datetime) else 'date' if isinstance(v,datetime.date) else 'object' if isinstance(v,dict) else 'array' if isinstance(v,list) else 'string' if isinstance(v,str) else 'unknown' for v in present)
    if kinds=={'integer','number'}: kinds={'number'}
    if len(kinds)>1 or 'unknown' in kinds: fail('Incompatible observed types for '+field)
    kind=next(iter(kinds),'string')
    col={'name':name(field),'logicalType':kind,'physicalType':{'string':'STRING','integer':'BIGINT','number':'DOUBLE','boolean':'BOOLEAN','timestamp':'TIMESTAMP','date':'DATE','object':'STRUCT','array':'ARRAY'}[kind], 'nullable':True,'evidence':'sampled-values','observedNulls':len(values)-len(present),'observedMissing':missing}
    if not present: col['warning']='All sampled values are null; type is unresolved, STRING is a draft placeholder'
    if kind=='integer' and any(v<-(2**63) or v>=2**63 for v in present): fail('Integer exceeds BIGINT for '+field)
    if kind=='number' and all(isinstance(v,(int,decimal.Decimal)) for v in present):
        numbers=[decimal.Decimal(v) for v in present]
        scale=max(max(0,-v.as_tuple().exponent) for v in numbers)
        integral=max(max(0,len(v.as_tuple().digits)+v.as_tuple().exponent) for v in numbers)
        precision=max(1,integral+scale)
        if precision>38: fail('Decimal precision exceeds 38')
        col['physicalType']=f'DECIMAL({precision},{scale})'
    if kind=='object': col['properties']=columns(present,depth+1)
    if kind=='array':
        elements=[v for row in present for v in row]
        if len(elements)>10000: fail('Array sample exceeds 10000 elements')
        col['items']=infer(elements,'item',depth=depth+1)
    return col

def columns(rows,depth=0):
    if not rows: fail('Empty dataset needs an explicit schema; no columns invented')
    if not all(isinstance(r,dict) for r in rows): fail('Expected object records')
    keys=list(dict.fromkeys(k for r in rows for k in r))
    if not keys or len(keys)>200: fail('Expected 1–200 columns')
    return [infer([r[k] for r in rows if k in r],k,sum(k not in r for r in rows),depth) for k in keys]

def arrow_column(f):
    import pyarrow as pa
    t=f.type
    if pa.types.is_dictionary(t): t=t.value_type
    if pa.types.is_struct(t):
        c={'logicalType':'object','physicalType':'STRUCT','properties':[arrow_column(x) for x in t]}
    elif pa.types.is_list(t) or pa.types.is_large_list(t):
        c={'logicalType':'array','physicalType':'ARRAY','items':arrow_column(pa.field('item',t.value_type))}
    elif pa.types.is_decimal(t):
        if t.precision>38 or t.scale<0 or t.scale>t.precision: fail('Unsupported Parquet decimal')
        c={'logicalType':'number','physicalType':f'DECIMAL({t.precision},{t.scale})'}
    elif pa.types.is_integer(t):
        if pa.types.is_uint64(t): fail('uint64 requires a reviewed decimal mapping')
        c={'logicalType':'integer','physicalType':'BIGINT'}
    elif pa.types.is_floating(t): c={'logicalType':'number','physicalType':'DOUBLE'}
    elif pa.types.is_boolean(t): c={'logicalType':'boolean','physicalType':'BOOLEAN'}
    elif pa.types.is_date(t): c={'logicalType':'date','physicalType':'DATE'}
    elif pa.types.is_timestamp(t):
        if t.tz: fail('Timezone-bearing Parquet timestamps need explicit conversion')
        c={'logicalType':'timestamp','physicalType':'TIMESTAMP'}
    elif pa.types.is_string(t) or pa.types.is_large_string(t): c={'logicalType':'string','physicalType':'STRING'}
    elif pa.types.is_binary(t) or pa.types.is_large_binary(t): c={'logicalType':'string','physicalType':'BINARY'}
    else: fail('Unsupported Parquet type')
    return dict(c,name=name(f.name),nullable=f.nullable,evidence='declared-parquet-schema')

def parse_file(raw,fmt,spec,limit):
    if fmt=='parquet':
        import pyarrow.parquet as pq
        p=pq.ParquetFile(io.BytesIO(raw))
        return [arrow_column(f) for f in p.schema_arrow], {'rowCount':p.metadata.num_rows,'sampledRows':0,'truncated':False}
    rows=[]
    if fmt in ['csv','tsv']:
        reader=csv.reader(io.StringIO(raw.decode('utf-8-sig'),newline=''),delimiter=',' if fmt=='csv' else '\t',strict=True)
        header=next(reader,[])
        if not header or len(header)!=len(set(header)): fail('Missing or duplicate delimited header')
        for h in header: name(h)
        for row in reader:
            if len(row)!=len(header): fail('Delimited row width differs from header')
            rows.append(dict(zip(header,[None if v=='' else v for v in row])))
            if len(rows)>limit: break
    elif fmt=='json':
        rows=parse_json(raw.decode('utf-8-sig'))
        if not isinstance(rows,list): fail('JSON discovery requires an explicit top-level record array')
    elif fmt=='jsonl':
        for line in raw.decode('utf-8-sig').splitlines():
            if line.strip(): rows.append(parse_json(line))
            if len(rows)>limit: break
    elif fmt=='xml':
        from defusedxml import ElementTree as ET
        root=ET.fromstring(raw,forbid_dtd=True,forbid_entities=True,forbid_external=True)
        for element in root:
            if element.tag!='row' or element.attrib: fail('XML requires simple root/row records')
            row={}
            for c in element:
                if c.tag in row or len(c) or set(c.attrib)-{'null'}: fail('Unsupported or duplicate XML field')
                if c.get('null') not in [None,'true']: fail('Unsupported XML null marker')
                if c.get('null')=='true' and c.text: fail('Null XML field contains text')
                row[name(c.tag)]=None if c.get('null')=='true' else (c.text or '')
            rows.append(row)
            if len(rows)>limit: break
    elif fmt=='xlsx':
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>2000 or sum(v.file_size for v in z.infolist())>32*1024*1024: fail('Workbook expanded size exceeds limit')
            if any('vbaproject' in v.filename.lower() or 'externallinks/' in v.filename.lower() for v in z.infolist()): fail('Macro/external-link workbook unsupported')
        wb=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            if spec.get('sheet') not in wb.sheetnames: fail('Select an existing explicit worksheet')
            iterator=wb[spec['sheet']].iter_rows()
            first=next(iterator,[])
            header=[c.value for c in first]
            if not header or len(header)!=len(set(header)): fail('Workbook header missing or duplicated')
            for h in header: name(h)
            for cells in iterator:
                if any(c.data_type=='f' for c in cells): fail('Formula cells require reviewed values; formulas are not evaluated')
                row=[c.value for c in cells]
                if not any(v is not None for v in row): continue
                rows.append(dict(zip(header,row)))
                if len(rows)>limit: break
        finally: wb.close()
    else: fail('Unsupported file format')
    sampled=rows[:limit]
    return columns(sampled), {'sampledRows':len(sampled),'truncated':len(rows)>limit}

class NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self,*args,**kwargs): fail('Storage redirects are not permitted')

def discover(config,read):
    if config.get('sourceKind')!='adls' or config.get('format') not in ['csv','tsv','json','jsonl','parquet','xml','xlsx']: fail('Unsupported source kind/format')
    if not isinstance(config.get('datasets'),list) or not 1<=len(config['datasets'])<=100 or not isinstance(config.get('sampleRows'),int) or not 1<=config['sampleRows']<=10000: fail('Invalid bounded source selection')
    names=set(); file_count=0
    for spec in config['datasets']:
        name(spec['name'])
        if spec['name'] in names: fail('Duplicate dataset name')
        names.add(spec['name'])
        if not isinstance(spec.get('paths'),list) or not 1<=len(spec['paths'])<=100: fail('Expected 1–100 paths per dataset')
        file_count+=len(spec['paths'])
        if len(set(spec['paths']))!=len(spec['paths']) or file_count>200: fail('Duplicate paths or more than 200 selected file reads')
    results=[]; total=0
    for spec in config['datasets']:
        found=[]; schemas=[]; sampled=0; truncated=False; row_count=0
        for path in spec['paths']:
            if not isinstance(path,str) or not re.fullmatch(r'[A-Za-z0-9_./= -]+',path) or any(p in ['', '.', '..'] for p in path.split('/')): fail('Unsafe source path')
            raw=read(path); total+=len(raw)
            if len(raw)>MAX_BYTES or total>MAX_TOTAL: fail('Source bytes exceed limit')
            cols,counts=parse_file(raw,config['format'],spec,config['sampleRows'])
            # Multiple files are accepted only with equal schemas; drift needs a new review.
            signature=[{k:v for k,v in c.items() if k not in ['observedNulls','observedMissing']} for c in cols]
            if schemas and signature!=schemas[0]: fail('Schema drift across explicitly selected files')
            schemas.append(signature)
            sampled+=counts['sampledRows']; truncated|=counts['truncated']; row_count+=counts.get('rowCount',0)
            found.append({'path':path,'sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw),**counts})
        partitions=spec.get('partitionColumns',[])
        for key in partitions:
            name(key)
            if any(c['name']==key for c in schemas[0]): fail('Partition column already exists in file schema')
            for path in spec['paths']:
                if not any(part.startswith(key+'=') and len(part)>len(key)+1 for part in path.split('/')[:-1]): fail('Missing declared path partition')
            schemas[0].append({'name':key,'logicalType':'string','physicalType':'STRING','nullable':True,'evidence':'declared-path-partition'})
        results.append({'name':spec['name'],'format':config['format'],'files':found,'sheet':spec.get('sheet'),'columns':schemas[0],'sampledRows':sampled,'truncated':truncated,**({'rowCount':row_count} if config['format']=='parquet' else {})})
    return {'apiVersion':'ingestron.file-metadata/v1','sourceId':config['sourceId'],'sourceKind':'adls','format':config['format'],'datasets':results}

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--config',default=str(pathlib.Path(__file__).with_name('discovery.json')))
    g=p.add_mutually_exclusive_group(required=True); g.add_argument('--local-root'); g.add_argument('--connections')
    p.add_argument('--out',required=True)
    args=p.parse_args(); out=pathlib.Path(args.out)
    if out.exists(): fail('Metadata destination exists')
    config=json.loads(pathlib.Path(args.config).read_text())
    if config.get('sourceKind')!='adls' or not 1<=len(config.get('datasets',[]))<=100 or not 1<=config.get('sampleRows',0)<=10000: fail('Invalid bounded discovery configuration')
    if args.local_root:
        root=pathlib.Path(args.local_root).resolve()
        def read(path):
            resolved=(root/path).resolve()
            if not resolved.is_relative_to(root): fail('Source escapes local root')
            with resolved.open('rb') as f: return f.read(MAX_BYTES+1)
    else:
        c=json.loads(pathlib.Path(args.connections).read_text())['storage']
        base=c['blobUrl'].rstrip('/'); token=c['sasToken'].lstrip('?')
        if not re.fullmatch(r'https://[a-z0-9]{3,24}\.blob\.core\.windows\.net/[a-z0-9-]{3,63}',base): fail('Expected an Azure Blob container HTTPS endpoint')
        opener=urllib.request.build_opener(NoRedirect())
        def read(path):
            try:
                url=base+'/'+urllib.parse.quote(path,safe='/=')+'?'+token
                with opener.open(url,timeout=60) as r: return r.read(MAX_BYTES+1)
            except Exception: fail('Blob read failed; check access, expiry and source path (credential URL suppressed)')
    result=discover(config,read)
    payload=json.dumps(result,indent=2)
    if len(payload.encode())>1900000: fail('Metadata exceeds CLI input bound; narrow source selection')
    with out.open('x',encoding='utf-8') as f: f.write(payload+'\n')
    print(json.dumps({'metadata':str(out),'datasets':len(result['datasets']),'format':config['format']}))
if __name__=='__main__':
    try: main()
    except Exception as exc:
        # Parser/HTTP exceptions can include source values or credential URLs.
        print(str(exc) if isinstance(exc,DiscoveryError) else 'Discovery failed; inspect format/settings and installed parser dependencies',file=sys.stderr)
        sys.exit(1)
